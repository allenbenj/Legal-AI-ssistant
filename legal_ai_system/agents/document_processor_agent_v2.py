from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any, Dict, List, Optional, Generic, TypeVar, cast

from ..core.base_agent import BaseAgent, AgentError
from ..core.detailed_logging import (
    LogCategory,
    detailed_log_function,
)
from ..utils.multimodal_types import (
    MultiModalDocument,
    AudioSegment,
    DepositionTranscript,
    StructuredForm,
    RedlineAnalysis,
)

InputT = TypeVar("InputT", bound=Path)
OutputT = TypeVar("OutputT", bound=MultiModalDocument)


class DocumentProcessorAgentV2(BaseAgent, Generic[InputT, OutputT]):
    """Enhanced document processor handling specialized legal content."""

    @detailed_log_function(LogCategory.AGENT)
    def __init__(self, service_container: Optional[Any] = None, **config: Any) -> None:
        super().__init__(service_container, name="DocumentProcessorAgentV2", agent_type="document_processing")
        if config:
            self.config.update(config)
        self.logger.info("DocumentProcessorAgentV2 initialized.")

    async def _process_task(self, task_data: Path, metadata: Dict[str, Any]) -> MultiModalDocument:
        """Trivial processing returning file contents as a ``MultiModalDocument``."""
        try:
            text = await asyncio.to_thread(Path(task_data).read_text, encoding="utf-8")
        except Exception as exc:  # pragma: no cover - runtime errors
            raise AgentError(f"Failed to read document: {exc}", self.name) from exc
        return MultiModalDocument(text_content=text)

    async def process_video_depositions(self, video_path: Path) -> DepositionTranscript:
        """Extract audio and transcribe a deposition video."""
        try:
            from moviepy.editor import VideoFileClip  # type: ignore
            import whisper  # type: ignore
        except Exception as exc:  # pragma: no cover - optional deps
            raise AgentError("moviepy and whisper are required for video processing", self.name) from exc

        audio_path = video_path.with_suffix(".wav")
        await asyncio.to_thread(self._extract_audio, video_path, audio_path)
        segments = await asyncio.to_thread(self._transcribe_audio, audio_path)
        return DepositionTranscript(segments=segments, metadata={"source": str(video_path)})

    def _extract_audio(self, video_path: Path, audio_path: Path) -> None:
        from moviepy.editor import VideoFileClip  # type: ignore

        clip = VideoFileClip(str(video_path))
        clip.audio.write_audiofile(str(audio_path), logger=None)

    def _transcribe_audio(self, audio_path: Path) -> List[AudioSegment]:
        import whisper  # type: ignore

        model = whisper.load_model("base")
        result = model.transcribe(str(audio_path))
        return [
            AudioSegment(
                start_ms=int(seg["start"] * 1000),
                end_ms=int(seg["end"] * 1000),
                transcript=seg["text"].strip(),
            )
            for seg in result.get("segments", [])
        ]

    async def process_legal_forms(self, form_path: Path) -> StructuredForm:
        """Detect fields within legal forms (PDF or Excel)."""
        if form_path.suffix.lower() == ".pdf":
            try:
                import pdfplumber  # type: ignore
            except Exception as exc:  # pragma: no cover - optional dep
                raise AgentError("pdfplumber required for PDF forms", self.name) from exc

            fields: Dict[str, str] = {}
            with pdfplumber.open(str(form_path)) as pdf:
                for page in pdf.pages:
                    for annot in page.annots or []:
                        name = annot.get("field_name") or annot.get("Subtype")
                        value = annot.get("field_value") or annot.get("Contents")
                        if name:
                            fields[str(name)] = str(value or "")
        elif form_path.suffix.lower() in {".xlsx", ".xls"}:
            try:
                import openpyxl  # type: ignore
            except Exception as exc:  # pragma: no cover
                raise AgentError("openpyxl required for Excel forms", self.name) from exc

            from openpyxl.worksheet.worksheet import Worksheet  # type: ignore

            wb = openpyxl.load_workbook(form_path, data_only=True)
            ws = cast(Worksheet, wb.active)
            fields = {
                str(r[0]): str(r[1])
                for r in ws.iter_rows(min_col=1, max_col=2, values_only=True)
                if r[0] is not None
            }
        else:
            raise AgentError("Unsupported form type", self.name)

        return StructuredForm(fields=fields, metadata={"source": str(form_path)})

    async def process_contract_redlines(self, doc_path: Path) -> RedlineAnalysis:
        """Parse tracked changes from a contract document."""
        try:
            import docx  # type: ignore
            from lxml import etree  # type: ignore
        except Exception as exc:  # pragma: no cover - optional deps
            raise AgentError("python-docx and lxml required for redline analysis", self.name) from exc

        doc = docx.Document(str(doc_path))
        root = etree.fromstring(doc._part.blob, parser=etree.XMLParser())
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        additions = [el.text for el in root.iter(f"{ns}ins") if el.text]
        deletions = [el.text for el in root.iter(f"{ns}del") if el.text]
        comments = [el.text for el in root.iter(f"{ns}comment") if el.text]

        return RedlineAnalysis(
            additions=additions,
            deletions=deletions,
            comments=comments,
            metadata={"source": str(doc_path)},
        )
